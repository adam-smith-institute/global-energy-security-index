"""Build the final Energy Security Index from the four step workbooks.

    pip install openpyxl             # once; needed to read the .xlsx workbooks
    python3 build_index.py

Reads   ../updated data 13:10:26/Step 1 ... Step 4 workbooks, and Step 0 to cross-check
Writes  ../Energy_Security_Index_2026_final.csv   one row per scored country, ranked

Rules, from the GESI 2026 methodology (draft of 12/10/2026):
  - Component scores are taken from the workbooks as calculated there; each is re-checked
    against its formula and any mismatch is reported.
  - A country enters the index only with all four component scores. There is no Step 3
    coverage threshold: Kiribati and Timor-Leste, below 90% partner coverage, are scored.
  - Overall = mean of the four components; rank 1 is the highest overall score.
  - The result must match the team's Step 0 overall ranking country for country.
Then run build_data.py to refresh the tracker.
"""
import csv, json, math, os, urllib.request
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "updated data 13:10:26")
OUT = os.path.join(HERE, "..", "Energy_Security_Index_2026_final.csv")
ISO_URL = "https://cdn.jsdelivr.net/npm/world-countries@5.1.0/countries.json"

# Plain display names for the trade data's official-style ones.
NAMES = {
    "USA": "United States", "Rep. of Korea": "South Korea", "Russian Federation": "Russia",
    "Brunei Darussalam": "Brunei", "Bolivia (Plurinational State of)": "Bolivia",
    "China, Macao SAR": "Macau", "China, Hong Kong SAR": "Hong Kong",
    "Bosnia Herzegovina": "Bosnia and Herzegovina", "United Rep. of Tanzania": "Tanzania",
    "Dominican Rep.": "Dominican Republic", "Rep. of Moldova": "Moldova",
    "Cayman Isds": "Cayman Islands", "Solomon Isds": "Solomon Islands",
    "Central African Rep.": "Central African Republic", "Dem. Rep. of the Congo": "DR Congo",
    "Congo": "Republic of the Congo", "Lao People's Dem. Rep.": "Laos",
    "Viet Nam": "Vietnam", "State of Palestine": "Palestine",
}


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v) else None


def clean(name):
    if "Ã" in name:  # UTF-8 text that was decoded as Latin-1 somewhere upstream
        try:
            name = name.encode("latin-1").decode("utf-8")
        except UnicodeError:
            pass
    return NAMES.get(name, name)


def rows_of(ws):
    return [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]


def src(name):
    return os.path.join(SRC, name)


problems = []

# Step 1: EIA generation, one block per country ending in a "Reliability Index" row.
s1, s1raw, cur = {}, {}, None
ws = openpyxl.load_workbook(src("Step 1. Oct 2.xlsx"), data_only=True, read_only=True).worksheets[0]
for r in ws.iter_rows(min_row=3, values_only=True):
    a, b, c = (list(r) + [None] * 3)[:3]
    if a is None and b and b != "Reliability Index":
        cur = b; s1raw[cur] = {}
    elif b == "Reliability Index":
        s1[cur] = num(c)
    elif cur and isinstance(a, str) and a.startswith("INTL"):
        s1raw[cur][b.split(" (")[0]] = num(c)

# Step 2: EIA consumption in five categories, HHI and score.
ws = openpyxl.load_workbook(src("Step 2. EIA_2024_Energy_Consumption_by_Country_5_Categories_with_HHI.xlsx"), data_only=True, read_only=True).worksheets[0]
s2, s2hhi = {}, {}
for v in rows_of(ws)[1:]:
    if v[0]:
        s2[v[0]], s2hhi[v[0]] = num(v[8]), num(v[7])

# Step 3: BACI import security.
wb3 = openpyxl.load_workbook(src("Step_3__BACI_2024_corrected.xlsx"), data_only=True, read_only=True)
calc = rows_of(wb3["Country Calc"])
s3 = {d["ISO3"]: d for d in (dict(zip(calc[0], r)) for r in calc[1:]) if d["ISO3"]}
par = {r[0]: r[1] for r in rows_of(wb3["Parameters"])}
BASE, WH, WD, WS = (par["Base import-dependence risk"], par["Supplier concentration weight (H)"],
                    par["Domestic instability weight (1-D)"], par["Supplier instability weight (1-S)"])

# Step 4: GDP per capita.
wb4 = openpyxl.load_workbook(src("Step_4__Final_corrected.xlsx"), data_only=True, read_only=True)
res4 = rows_of(wb4["Step 4 Results"])
s4 = {r[2]: dict(zip(res4[0], r)) for r in res4[1:]}
p4 = {r[0]: r[1] for r in rows_of(wb4["Parameters"])}
LNMIN, LNMAX = math.log(p4["GDPpc minimum"]), math.log(p4["GDPpc maximum"])

# Re-check every component score against its formula.
for name, raw in s1raw.items():
    g = raw.get("Generation")
    if s1.get(name) is None or not g:
        continue
    exp1 = 100 * (g - sum(raw.get(k) or 0 for k in ("Wind", "Solar", "Tide and wave"))) / g
    if abs(exp1 - s1[name]) > 1e-6:
        problems.append(f"Step 1 {name}: workbook {s1[name]:.4f}, formula {exp1:.4f}")

included, excluded = [], []
for iso, d in s3.items():
    name, eia = clean(d["Country"]), d["EIA Source Country"]
    v = dict(s1=s1.get(eia), s2=s2.get(eia), s3=num(d["Step 3 Score (0-100)"]), s4=(s4.get(iso) or {}).get("Step 4 Score (0-100)"))
    missing = [k for k, x in v.items() if x is None]
    if missing:
        excluded.append((iso, name, "missing " + ", ".join(missing))); continue
    if s2hhi.get(eia) is not None and abs(100 * (1 - s2hhi[eia]) / 0.8 - v["s2"]) > 1e-6:
        problems.append(f"Step 2 {name}: score does not match its HHI")
    I = d["I: Net Import Dependence"]
    exp3 = 100.0 if I == 0 else 100 * (1 - I * (BASE + WH * d["H: Supplier HHI"] + WD * (1 - d["D (0-1)"]) + WS * (1 - d["S: Supplier Stability"])))
    if abs(exp3 - v["s3"]) > 1e-6:
        problems.append(f"Step 3 {name}: workbook {v['s3']:.4f}, formula {exp3:.4f}")
    gdp = s4[iso]["GDP per Capita (constant 2015 US$)"]
    if abs(100 * (math.log(gdp) - LNMIN) / (LNMAX - LNMIN) - v["s4"]) > 1e-6:
        problems.append(f"Step 4 {name}: score does not match its GDP per capita")
    pc, wc = d["Identifiable Partner Coverage"] or 0, d["Supplier WGI Coverage"] or 0
    included.append(dict(
        iso3=iso, country=name, overall=sum(v.values()) / 4, **v,
        net_import_dependence=I, partner_hhi=d["H: Supplier HHI"], domestic_wgi=d["Domestic WGI (0-100)"],
        supplier_wgi=(100 * d["S: Supplier Stability"]) if d["S: Supplier Stability"] is not None else "",
        partner_coverage_pct=100 * pc, supplier_wgi_coverage_pct=100 * wc, supplier_count=d["Supplier Count"],
        gdp_per_capita_2015usd=gdp))

included.sort(key=lambda r: -r["overall"])
for r in included:
    r["rank"] = 1 + sum(1 for o in included if o["overall"] > r["overall"] + 1e-12)

iso_numeric = {c["cca3"]: c.get("ccn3") for c in json.load(urllib.request.urlopen(ISO_URL))}
cols = ["rank", "iso3", "iso_numeric", "country", "overall", "s1", "s2", "s3", "s4",
        "net_import_dependence", "partner_hhi", "domestic_wgi", "supplier_wgi",
        "partner_coverage_pct", "supplier_wgi_coverage_pct", "supplier_count", "gdp_per_capita_2015usd"]
with open(OUT, "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=cols)
    w.writeheader()
    for r in included:
        w.writerow({k: (r.get(k) if k != "iso_numeric" else iso_numeric.get(r["iso3"], "")) for k in cols})

print(f"{len(included)} countries scored -> {os.path.normpath(OUT)}")
print(f"{len(excluded)} Step 3 economies excluded for a missing component")
print("scored with under 90% partner coverage:",
      [f"{r['country']} ({r['partner_coverage_pct']:.1f}%)" for r in included if r["partner_coverage_pct"] < 90])

# Cross-check against the team's own final ranking.
ov = rows_of(openpyxl.load_workbook(src("Step_0__Overall_Ranking_2024_corrected.xlsx"), data_only=True, read_only=True)["Overall Ranking"])
step0 = {d["ISO3"]: d for d in (dict(zip(ov[0], r)) for r in ov[1:])}
mine = {r["iso3"]: r for r in included}
if set(step0) != set(mine):
    problems.append(f"Step 0 universe differs: only in Step 0 {sorted(set(step0) - set(mine))}, only here {sorted(set(mine) - set(step0))}")
for iso in set(step0) & set(mine):
    a, b = step0[iso], mine[iso]
    pairs = [("Overall Score", "overall")] + [(f"Step {i} Score", f"s{i}") for i in range(1, 5)]
    bad = [k for k, m in pairs if abs(a[k] - b[m]) > 1e-4] + (["rank"] if int(a["Overall Rank"]) != b["rank"] else [])
    if bad:
        problems.append(f"Step 0 {b['country']}: differs on {', '.join(bad)}")
print("formula checks:", "all match" if not problems else f"{len(problems)} mismatches")
for p in problems:
    print("  ", p)
